"""把试卷写成 A4 docx。学生卷函数不读取答案字段。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from cn_exam.models import Paper, PaperMeta, SafeQuestion, SafeSection, to_safe_section
from cn_exam.tianzi import render_tianzi_png

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
CN_ORDINAL = "一二三四五六七八九十"
TIANZI_HEIGHT = Mm(12)


def render_student(paper: Paper, path: Path) -> None:
    """写出学生卷。先丢掉答案字段，再排版。"""
    safe_sections = [to_safe_section(section) for section in paper.sections]
    _write_student(paper.meta, safe_sections, list(paper.warnings), path)


def render_answer(paper: Paper, path: Path) -> None:
    """写出答案卷。答案行统一以“答案：”开头，便于测试检查学生卷。"""
    document = _new_document()
    meta = paper.meta
    _add_running_header(document, meta)
    _add_title(document, "小学语文 · 参考答案", meta)
    _add_center(
        document,
        f"配套学生卷：{meta.scope_label}　　生成种子：{meta.seed}　　满分 {meta.total_points} 分",
        10.5,
        MUTED,
    )
    _add_center(document, "本答案只对应同一种子生成的学生卷，不要印在学生卷上。", 10.5, ACCENT)
    for index, section in enumerate(paper.sections, start=1):
        heading = f"{_ordinal(index)}、{section.title}（{section.points}分）"
        _add_heading(document, heading)
        for question in section.questions:
            stem = question.visible_parts[0] if question.visible_parts else ""
            _add_body(document, f"{question.number}.（{question.points}分）{stem}")
            for extra in question.visible_parts[1:]:
                _add_body(document, extra, size=10.5, color=MUTED)
            for option_line in _option_lines(question.options):
                _add_body(document, option_line, size=10.5, color=MUTED)
            for line_index, line in enumerate(question.answer.split("\n")):
                prefix = "答案：" if line_index == 0 else ""
                _add_body(document, prefix + line, bold=True, color=ACCENT)
            if question.explanation:
                _add_body(document, question.explanation, size=10.5, color=MUTED)
    _add_notes(document, meta, paper.warnings)
    _save(document, path)


def _write_student(
    meta: PaperMeta,
    sections: list[SafeSection],
    warnings: list[str],
    path: Path,
) -> None:
    for section in sections:
        if not isinstance(section, SafeSection):
            raise TypeError("学生卷只能接收不含答案的 SafeSection")
        for question in section.questions:
            if not isinstance(question, SafeQuestion):
                raise TypeError("学生卷只能接收不含答案的 SafeQuestion")
    document = _new_document()
    _add_running_header(document, meta)
    _add_title(document, f"小学语文 · {meta.book_name} · {meta.scope_label}", meta)
    _add_info_table(document, meta)
    _add_body(
        document,
        f"满分 {meta.total_points} 分　　建议用时 {meta.duration_minutes} 分钟",
        size=10.5,
        color=MUTED,
    )
    _add_body(document, f"范围：{meta.scope_detail}", size=10.5, color=MUTED)
    for index, section in enumerate(sections, start=1):
        _add_heading(document, f"{_ordinal(index)}、{section.title}（{section.points}分）")
        instruction = section.instruction
        if _uniform_points(section) and section.id != "writing":
            instruction += f"每小题{section.questions[0].points}分。"
        _add_body(document, instruction, size=10.5, color=MUTED)
        for preamble_index, paragraph in enumerate(section.preamble):
            if section.id == "reading" and preamble_index == 0:
                _add_body(document, paragraph, bold=True, align="center")
            elif section.id == "reading":
                _add_body(document, "　　" + paragraph)
            else:
                _add_body(document, paragraph)
        if section.id == "pinyin":
            _add_pinyin_table(document, section)
            continue
        for question in section.questions:
            _add_question(document, question, show_points=not _uniform_points(section))
    _add_notes(document, meta, warnings)
    _save(document, path)


def _add_question(document: Document, question: SafeQuestion, show_points: bool) -> None:
    points = f"（{question.points}分）" if show_points or question.kind == "classical" else ""
    if question.kind in {"judge", "single"}:
        first = f"（　　）{question.number}. {points}{question.visible_parts[0]}"
        _add_body(document, first)
        for extra in question.visible_parts[1:]:
            _add_body(document, extra)
        for option_line in _option_lines(question.options):
            _add_body(document, option_line)
        return
    if question.kind == "writing":
        for paragraph in question.visible_parts:
            _add_body(document, paragraph)
        for _ in range(question.writing_lines):
            _add_writing_line(document)
        return
    first = f"{question.number}.{points}{question.visible_parts[0]}" if question.visible_parts else ""
    if first:
        _add_body(document, first)
    for extra in question.visible_parts[1:]:
        _add_body(document, extra)
    for grid in question.blank_grids:
        _add_tianzi(document, grid)
    if question.kind == "reading":
        # 题干和第一条横线绑在一起，避免题干落在页末、横线掉到下一页。
        if question.visible_parts:
            document.paragraphs[-1].paragraph_format.keep_with_next = True
        _add_body(document, "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿", keep_with_next=True)
        _add_body(document, "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿")


def _add_pinyin_table(document: Document, section: SafeSection) -> None:
    rows = (len(section.questions) + 1) // 2
    table = document.add_table(rows=rows, cols=2)
    _set_table_widths(table, [Cm(8.6), Cm(8.6)])
    _clear_table_borders(table)
    show_points = not _uniform_points(section)
    for index, question in enumerate(section.questions):
        cell = table.rows[index // 2].cells[index % 2]
        _set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        points = f"（{question.points}分）" if show_points else ""
        pinyin = question.visible_parts[0] if question.visible_parts else ""
        _add_run(paragraph, f"{question.number}. {points}{pinyin}", 12, False, INK, "宋体")
        grid = question.blank_grids[0] if question.blank_grids else 2
        image_paragraph = cell.add_paragraph()
        image_paragraph.paragraph_format.space_before = Pt(0)
        image_paragraph.paragraph_format.space_after = Pt(4)
        _add_tianzi_run(image_paragraph, grid)
    _tighten_after_table(document)


def _add_tianzi(document: Document, count: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    _add_tianzi_run(paragraph, count)


def _add_tianzi_run(paragraph, count: int) -> None:
    stream = BytesIO(render_tianzi_png(count))
    stream.seek(0)
    run = paragraph.add_run()
    run.add_picture(stream, height=TIANZI_HEIGHT)


def _add_writing_line(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    _add_run(paragraph, "\u00A0", 12, False, INK, "宋体")
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B7C3D0")
    borders.append(bottom)
    paragraph_properties.append(borders)


def _new_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    section.different_first_page_header_footer = True
    _set_style_font(document)
    _add_page_footer(section.footer)
    _add_page_footer(section.first_page_footer)
    return document


def _add_running_header(document: Document, meta: PaperMeta) -> None:
    """从第 2 页开始显示试卷名称，首页用卷头大标题。"""
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(paragraph, f"{meta.book_name} · {meta.scope_label}", 9, False, MUTED, "宋体")


def _add_title(document: Document, text: str, meta: PaperMeta) -> None:
    _add_center(document, text, 18, ACCENT, bold=True, font="黑体", space_after=2)
    subtitle = f"{meta.edition_name}  ·  {meta.series_label}"
    _add_center(document, subtitle, 11, MUTED, space_after=6)


def _add_info_table(document: Document, meta: PaperMeta) -> None:
    table = document.add_table(rows=1, cols=4)
    _set_table_widths(table, [Cm(4.3), Cm(4.3), Cm(4.3), Cm(4.3)])
    labels = ("姓名：", "班级：", "学号：", "得分：")
    for cell, label in zip(table.rows[0].cells, labels):
        _shade_cell(cell, "F4F7FB")
        _set_cell_borders(cell, "D5DDE8")
        _set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_run(paragraph, f"{label}＿＿＿＿", 11, False, INK, "宋体")
    _tighten_after_table(document)
    # 得分栏用卷头信息，不预填分数。meta 仅用于保证版式与试卷对应。
    _ = meta


def _add_notes(document: Document, meta: PaperMeta, warnings: list[str]) -> None:
    if meta.paper_note:
        _add_body(document, meta.paper_note, size=9, color=MUTED)
    for warning in warnings:
        _add_body(document, warning, size=9, color=MUTED)


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    _add_run(paragraph, text, 14, True, ACCENT, "黑体")
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    borders.append(bottom)
    properties.append(borders)


def _add_center(
    document: Document,
    text: str,
    size: float,
    color: RGBColor,
    bold: bool = False,
    font: str = "宋体",
    space_after: float = 2,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    _add_run(paragraph, text, size, bold, color, font)


def _add_body(
    document: Document,
    text: str,
    size: float = 12,
    color: RGBColor = INK,
    bold: bool = False,
    align: str = "left",
    keep_with_next: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = keep_with_next
    _add_run(paragraph, text, size, bold, color, "宋体")


def _add_run(paragraph, text: str, size: float, bold: bool, color: RGBColor, font: str) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Times New Roman"
    fonts = run._element.rPr.rFonts
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), font)
    fonts.set(qn("w:cs"), "Times New Roman")


def _option_lines(options: tuple[str, ...] | list[str]) -> list[str]:
    """没有选项时不占空行。选项较短时排在同一行。"""
    if not options:
        return []
    text = _format_options(options)
    return [line for line in text.split("\n") if line]


def _format_options(options: tuple[str, ...] | list[str]) -> str:
    labeled = [f"{chr(ord('A') + index)}. {option}" for index, option in enumerate(options)]
    if sum(len(item) for item in labeled) <= 30:
        return "　　".join(labeled)
    return "\n".join(labeled)


def _uniform_points(section: SafeSection) -> bool:
    points = {question.points for question in section.questions}
    return len(points) == 1


def _ordinal(index: int) -> str:
    if not 1 <= index <= len(CN_ORDINAL):
        raise ValueError(f"题型序号超出范围：{index}")
    return CN_ORDINAL[index - 1]


def _set_style_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = INK
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


def _add_page_footer(footer) -> None:
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(paragraph, "第 ", 9, False, MUTED, "宋体")
    _add_page_field(paragraph)
    _add_run(paragraph, " 页", 9, False, MUTED, "宋体")


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    run.font.size = Pt(9)


def _shade_cell(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    shading.set(qn("w:val"), "clear")
    properties.append(shading)


def _set_cell_borders(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    properties.append(borders)


def _set_cell_margins(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", "60"), ("left", "80"), ("bottom", "60"), ("right", "80")):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)


def _clear_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")
        borders.append(element)
    properties.append(borders)


def _set_table_widths(table, widths) -> None:
    """按厘米锁定列宽，避免田字格把表格撑出页面。"""
    table.autofit = False
    table.allow_autofit = False
    widths_dxa = [int(width.twips) for width in widths]
    total = sum(widths_dxa)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        properties.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _tighten_after_table(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)


def _save(document: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
